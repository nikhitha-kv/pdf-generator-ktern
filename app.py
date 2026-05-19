import os
import json
import uuid
import base64
import requests
import re
from datetime import datetime
from flask import Flask, render_template, request, jsonify, send_from_directory, url_for
from dotenv import load_dotenv
from werkzeug.utils import secure_filename
import pdfplumber
import docx
import pandas as pd

# ReportLab imports
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch, cm
from reportlab.lib import colors
from reportlab.platypus import BaseDocTemplate, PageTemplate, Frame, Paragraph, Spacer, Image, Table, TableStyle, PageBreak
from reportlab.platypus.tableofcontents import TableOfContents

# Load environment variables
load_dotenv(override=True)

app = Flask(__name__, template_folder='templates', static_folder='static')

# Constants and Configuration
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
PDF_DIR = os.path.join(BASE_DIR, 'generated_pdfs')
ASSETS_DIR = os.path.join(BASE_DIR, 'assets')
os.makedirs(PDF_DIR, exist_ok=True)
os.makedirs(ASSETS_DIR, exist_ok=True)

AI_MODEL = os.getenv('AI_MODEL', 'google/gemini-2.5-pro')

def fetch_pexels_image(query):
    pexels_key = os.getenv('PEXELS_API_KEY', '')
    if not pexels_key:
        print("Warning: Pexels API key not set. Skipping image fetch.")
        return None
        
    url = "https://api.pexels.com/v1/search"
    headers = {"Authorization": pexels_key}
    params = {"query": query + " electronic hardware white background", "per_page": 1, "orientation": "landscape"}
    
    try:
        response = requests.get(url, headers=headers, params=params)
        response.raise_for_status()
        data = response.json()
        
        if data.get('photos') and len(data['photos']) > 0:
            img_url = data['photos'][0]['src']['medium']
            img_response = requests.get(img_url)
            
            filename = f"pexels_{uuid.uuid4().hex}.jpg"
            filepath = os.path.join(ASSETS_DIR, filename)
            
            with open(filepath, 'wb') as f:
                f.write(img_response.content)
            return filepath
    except Exception as e:
        print(f"Error fetching image from Pexels for query '{query}': {e}")
    return None

def fetch_mermaid_flowchart(mermaid_code):
    try:
        # Mermaid.ink requires base64 encoded graph definition
        encoded = base64.urlsafe_b64encode(mermaid_code.encode('utf-8')).decode('utf-8')
        url = f"https://mermaid.ink/img/{encoded}"
        
        response = requests.get(url)
        response.raise_for_status()
        
        filename = f"flowchart_{uuid.uuid4().hex}.png"
        filepath = os.path.join(ASSETS_DIR, filename)
        
        with open(filepath, 'wb') as f:
            f.write(response.content)
        return filepath
    except Exception as e:
        print(f"Error generating flowchart: {e}")
    return None

def generate_report_data(topic, requirements):
    openrouter_key = os.getenv('OPENROUTER_API_KEY', '')
    print(f"DEBUG KEY: {openrouter_key[:10]}...{openrouter_key[-5:]}", flush=True)
    if not openrouter_key:
        raise ValueError("OpenRouter API key is missing in .env")
    prompt = f"""
    You are an elite enterprise architect, program manager, and lead SAP consultant by KaarTech. Generate a comprehensive, professional enterprise document (e.g. SAP Migration Proposal, BRD, Project Charter) based on the following:
    Topic: {topic}
    Requirements/Context: {requirements}

    CRITICAL TONE RULES:
    1. Be concise, factual, structured, and business-focused.
    2. DO NOT use AI-sounding marketing clichés. 
    3. ABSOLUTELY BANNED PHRASES: "revolutionary", "industry leading", "proprietary framework", "guaranteed", "accelerate transformation journey", "next generation", "world class".
    You MUST output the result in STRICT JSON format with NO markdown wrapping.
    
    The JSON MUST have this structure:
    {{
      "title": "Full Document Title here",
      "client": "Client Name or Enterprise",
      "date": "YYYY-MM-DD",
      "sections": [
        {{ "id": "section-id-1", "title": "1. Section Title 1", "content": "HTML content here..." }},
        {{ "id": "section-id-2", "title": "2. Section Title 2", "content": "HTML content here..." }}
      ]
    }}

    CRITICAL STRUCTURE RULES:
    1. Generate between 4 to 12 logical sections tailored specifically to the document type/topic.
    2. Define relevant sections that cover all aspects of the requested topic (e.g. if the topic is a cutover plan, include phases, tasks, checklist, rollback path. If it's a security role matrix, include roles, permissions, GRC mapping. If it's an enterprise overview like "kaartech overview", include company vision, offerings, SAP expertise).
    
    IMPORTANT CONTENT RULES:
    1. The 'content' field MUST contain well-formatted, rich HTML (<p>, <ul>, <li>, <strong>, <table class="doc-table">).
    2. NEVER generate the timeline as paragraphs. It MUST be an HTML table with Phase | Duration | Activities | Owner | Dependency.
    3. NEVER generate risks as paragraphs. It MUST be an HTML table with Risk | Impact | Severity | Mitigation. Use <span class="badge high">High</span>, <span class="badge medium">Medium</span>, <span class="badge low">Low</span> for the Severity column.
    4. For any architecture, processes, or diagrams, DO NOT use simple bullet points or placeholders. You MUST generate a valid, standard Mermaid flowchart inside a `<div class="mermaid">` tag, using standard layout syntax (e.g., `graph TD` or `graph LR` with styled nodes).
    5. Ensure the JSON is valid and escaped correctly.
    """

    headers = {
        "Authorization": f"Bearer {openrouter_key}",
        "Content-Type": "application/json",
        "HTTP-Referer": "http://localhost:5000",
        "X-Title": "KTern Enterprise AI Document Generator"
    }
    
    payload = {
        "model": AI_MODEL,
        "max_tokens": 6000,
        "messages": [
            {"role": "system", "content": "You are a JSON-generating enterprise SAP consulting assistant. Output ONLY valid JSON."},
            {"role": "user", "content": prompt}
        ],
        "temperature": 0.3
    }
    
    models_to_try = [
        AI_MODEL,
        "poolside/laguna-m.1:free",
        "google/gemma-4-26b-a4b-it:free",
        "nvidia/nemotron-3-nano-30b-a3b:free",
        "openai/gpt-oss-120b:free"
    ]
    
    response = None
    last_error = None
    
    for model_name in models_to_try:
        payload["model"] = model_name
        print(f"Attempting generation with model: {model_name}...", flush=True)
        try:
            r = requests.post("https://openrouter.ai/api/v1/chat/completions", headers=headers, json=payload, timeout=60)
            r.raise_for_status()
            response = r
            print(f"Successfully generated with {model_name}!", flush=True)
            break
        except Exception as e:
            print(f"Model {model_name} failed: {e}. Trying next fallback...", flush=True)
            last_error = e
            
    if not response:
        raise RuntimeError(f"All AI models failed to generate content. Last error: {last_error}")
    
    response_text = response.json()['choices'][0]['message']['content']
    
    # Strip markdown if the AI mistakenly included it
    response_text = re.sub(r'^```json\s*', '', response_text)
    response_text = re.sub(r'^```\s*', '', response_text)
    response_text = re.sub(r'\s*```$', '', response_text)
    
    try:
        return json.loads(response_text)
    except json.JSONDecodeError as e:
        print("Failed to decode JSON:", response_text[:200], "...")
        raise ValueError("AI response was not valid JSON.")





def clean_html_for_reportlab(html_text):
    if not html_text:
        return ""
    
    # Normalize common problematic unicode characters for ReportLab
    html_text = html_text.replace('\u2011', '-')
    html_text = html_text.replace('\u2013', '-')
    html_text = html_text.replace('\u2014', '--')
    html_text = html_text.replace('\u201c', '"').replace('\u201d', '"')
    html_text = html_text.replace('\u2018', "'").replace('\u2019', "'")
    
    # Convert strong/em to b/i
    html_text = html_text.replace('<strong>', '<b>').replace('</strong>', '</b>')
    html_text = html_text.replace('<em>', '<i>').replace('</em>', '</i>')
    
    # Convert headings to bold + break
    for i in range(1, 7):
        html_text = html_text.replace(f'<h{i}>', '<b>').replace(f'</h{i}>', '</b><br/>')
        
    html_text = html_text.replace('<br>', '<br/>')
    
    # Strip block elements and any other tags ReportLab Paragraph doesn't support
    def tag_replacer(match):
        full_tag = match.group(0)
        tag_name = match.group(1).lower()
        if tag_name in ['b', 'i', 'u', 'font', 'a']:
            return full_tag
        elif tag_name in ['br']:
            return '<br/>'
        return '' # strip other tags
        
    html_text = re.sub(r'</?([a-zA-Z0-9]+)(?:\s+[^>]*)?>', tag_replacer, html_text)
    
    # Escape ampersands but avoid double-escaping entities
    html_text = re.sub(r'&(?!amp;|lt;|gt;|quot;|apos;|nbsp;)', '&amp;', html_text)
    
    # Balance remaining inline tags (b, i, u, font, a)
    for tag in ['b', 'i', 'u', 'font', 'a']:
        open_count = len(re.findall(f'<{tag}(?:\\s+[^>]*)?>', html_text, re.IGNORECASE))
        close_count = len(re.findall(f'</{tag}>', html_text, re.IGNORECASE))
        
        if open_count > close_count:
            html_text += f'</{tag}>' * (open_count - close_count)
        elif close_count > open_count:
            html_text = f'<{tag}>' * (close_count - open_count) + html_text
            
    # Remove excess breaks
    html_text = re.sub(r'(<br/>\s*){3,}', '<br/><br/>', html_text)
    html_text = html_text.strip()
    html_text = re.sub(r'^(<br/>\s*)+', '', html_text)
    html_text = re.sub(r'(\s*<br/>)+$', '', html_text)
    html_text = html_text.strip()
    
    return html_text


def parse_html_to_story(html_content, styles, story, brand_navy, brand_red, border_grey, A4):
    content = (html_content or '').strip()
    if not content:
        return

    # Regex to extract structured block elements sequentially
    pattern = re.compile(
        r'(<table.*?>.*?</table>|'
        r'<div class=["\']timeline-list["\'].*?>.*?</div>|'
        r'<div class=["\']mermaid["\'].*?>.*?</div>|'
        r'<ul.*?>.*?</ul>|'
        r'<ol.*?>.*?</ol>|'
        r'<p.*?>.*?</p>|'
        r'<div class=["\']risk-item["\'].*?>.*?</div>|'
        r'<div class=["\']diagram-container["\'].*?>.*?</div>)',
        re.DOTALL | re.IGNORECASE
    )

    matches = list(pattern.finditer(content))

    if not matches:
        story.append(Paragraph(clean_html_for_reportlab(content), styles['CorpBody']))
        story.append(Spacer(1, 6))
        return

    last_end = 0
    for match in matches:
        between_text = content[last_end:match.start()].strip()
        if between_text:
            clean_text = clean_html_for_reportlab(between_text)
            if clean_text:
                story.append(Paragraph(clean_text, styles['CorpBody']))
                story.append(Spacer(1, 6))

        block_html = match.group(1)
        last_end = match.end()

        # Handle specific blocks
        if block_html.lower().startswith('<table'):
            try:
                headers = re.findall(r'<th.*?>(.*?)</th>', block_html, re.DOTALL | re.IGNORECASE)
                rows_html = re.findall(r'<tr.*?>(.*?)</tr>', block_html, re.DOTALL | re.IGNORECASE)

                table_data = []
                if headers:
                    table_data.append([Paragraph(clean_html_for_reportlab(h), styles['CorpTableHeader']) for h in headers])

                for r_html in rows_html:
                    if '<th' in r_html.lower() and not '<td' in r_html.lower():
                        continue
                    cols = re.findall(r'<td.*?>(.*?)</td>', r_html, re.DOTALL | re.IGNORECASE)
                    if cols:
                        table_data.append([Paragraph(clean_html_for_reportlab(c), styles['CorpTableText']) for c in cols])

                if table_data:
                    num_cols = len(table_data[0])
                    col_width = (A4[0] - 2 * inch) / num_cols
                    t = Table(table_data, colWidths=[col_width] * num_cols)
                    t_style = [
                        ('BACKGROUND', (0, 0), (-1, 0), brand_navy),
                        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                        ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
                        ('TOPPADDING', (0, 0), (-1, 0), 6),
                        ('GRID', (0, 0), (-1, -1), 0.5, border_grey),
                    ]
                    for r_idx in range(1, len(table_data)):
                        if r_idx % 2 == 0:
                            t_style.append(('BACKGROUND', (0, r_idx), (-1, r_idx), colors.HexColor('#f9fafb')))
                    t.setStyle(TableStyle(t_style))
                    story.append(t)
                    story.append(Spacer(1, 8))
            except Exception as e:
                print(f"PDF Table Parse Error: {e}")
                story.append(Paragraph(clean_html_for_reportlab(block_html), styles['CorpBody']))
                story.append(Spacer(1, 6))

        elif 'timeline-list' in block_html.lower():
            try:
                phases = re.findall(r"<div class=['\"]timeline-phase['\"]>(.*?)</div>", block_html, re.DOTALL | re.IGNORECASE)
                durations = re.findall(r"<div class=['\"]timeline-duration['\"]>(.*?)</div>", block_html, re.DOTALL | re.IGNORECASE)
                descs = re.findall(r"<div class=['\"]timeline-desc['\"]>(.*?)</div>", block_html, re.DOTALL | re.IGNORECASE)

                timeline_table_data = [[
                    Paragraph("<b>Migration Phase</b>", styles['CorpTableHeader']),
                    Paragraph("<b>Duration</b>", styles['CorpTableHeader']),
                    Paragraph("<b>Details & Objectives</b>", styles['CorpTableHeader'])
                ]]

                for idx in range(min(len(phases), len(durations), len(descs))):
                    timeline_table_data.append([
                        Paragraph(clean_html_for_reportlab(phases[idx]), styles['CorpTableText']),
                        Paragraph(f"<font color='#c8102e'><b>{clean_html_for_reportlab(durations[idx])}</b></font>", styles['CorpTableText']),
                        Paragraph(clean_html_for_reportlab(descs[idx]), styles['CorpTableText'])
                    ])

                t = Table(timeline_table_data, colWidths=[1.8 * inch, 1.2 * inch, 3.5 * inch])
                t_style = [
                    ('BACKGROUND', (0, 0), (-1, 0), brand_navy),
                    ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                    ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                    ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
                    ('TOPPADDING', (0, 0), (-1, 0), 6),
                    ('GRID', (0, 0), (-1, -1), 0.5, border_grey),
                ]
                for r_idx in range(1, len(timeline_table_data)):
                    if r_idx % 2 == 0:
                        t_style.append(('BACKGROUND', (0, r_idx), (-1, r_idx), colors.HexColor('#f9fafb')))
                t.setStyle(TableStyle(t_style))
                story.append(t)
                story.append(Spacer(1, 8))
            except Exception as e:
                print(f"PDF Timeline Parse Error: {e}")
                story.append(Paragraph(clean_html_for_reportlab(block_html), styles['CorpBody']))
                story.append(Spacer(1, 6))

        elif 'mermaid' in block_html.lower() or 'diagram-container' in block_html.lower():
            try:
                # Extract code inside mermaid div or fallback to diagram-container code
                mermaid_code = re.search(r'<div class=["\']mermaid["\']>(.*?)</div>', block_html, re.DOTALL | re.IGNORECASE)
                code = ""
                if mermaid_code:
                    code = mermaid_code.group(1).strip()
                else:
                    code = """graph TD
    classDef client fill:#f8fafc,stroke:#475569,stroke-width:1.5px
    classDef web fill:#eff6ff,stroke:#2563eb,stroke-width:1.5px
    classDef app fill:#fef3c7,stroke:#d97706,stroke-width:2px
    classDef db fill:#fff1f2,stroke:#e11d48,stroke-width:1.5px

    Client["Client / Mobile UX<br/>SAP Fiori Launchpad"]:::client
    Gateway["SAP Gateway Server<br/>OData / HTTPS"]:::web
    AppServer["SAP S/4HANA Core<br/>ABAP Application Server"]:::app
    DBServer["SAP HANA DB Layer<br/>In-Memory Database"]:::db

    Client -->|HTTPS Port 443| Gateway
    Gateway -->|RFC Connection| AppServer
    AppServer -->|SQL In-Memory Access| DBServer"""

                # Clean entities
                code = code.replace("&gt;", ">").replace("&lt;", "<").replace("&amp;", "&")

                # Guarantee double quotes around labels containing slashes or spaces
                lines = code.split('\n')
                fixed_lines = []
                for line in lines:
                    m = re.match(r'^(\s*)(\w+)(\[[^"\].]+\]|\([^"\).]+\)|\{[^"\}.]+\})(.*)$', line)
                    if m:
                        indent, node_id, text_bracket, rest = m.groups()
                        bracket_open = text_bracket[0]
                        bracket_close = text_bracket[-1]
                        text_inside = text_bracket[1:-1]
                        if not text_inside.startswith('"'):
                            line = f'{indent}{node_id}{bracket_open}"{text_inside}"{bracket_close}{rest}'
                    fixed_lines.append(line)
                code = '\n'.join(fixed_lines)

                img_path = fetch_mermaid_flowchart(code)
                if img_path:
                    img = Image(img_path)
                    img.drawHeight = 3.5 * inch
                    img.drawWidth = 5.0 * inch
                    img.hAlign = 'CENTER'
                    story.append(img)
                    story.append(Spacer(1, 10))
                else:
                    story.append(Paragraph("Architecture Diagram Generation Failed.", styles['CorpBody']))
            except Exception as e:
                print(f"PDF Diagram Generation Error: {e}")
                story.append(Paragraph("Architecture Diagram Failed.", styles['CorpBody']))

        elif block_html.lower().startswith('<ul') or block_html.lower().startswith('<ol'):
            try:
                bullets = re.findall(r'<li.*?>(.*?)</li>', block_html, re.DOTALL | re.IGNORECASE)
                for b in bullets:
                    story.append(Paragraph(f"&bull; {clean_html_for_reportlab(b)}", styles['CorpBullet']))
                story.append(Spacer(1, 4))
            except Exception as e:
                story.append(Paragraph(clean_html_for_reportlab(block_html), styles['CorpBody']))
                story.append(Spacer(1, 6))

        elif 'risk-item' in block_html.lower():
            try:
                titles = re.findall(r"<div class=['\"]risk-title['\"]>(.*?)</div>", block_html, re.DOTALL | re.IGNORECASE)
                descs = re.findall(r"<div class=['\"]risk-desc['\"]>(.*?)</div>", block_html, re.DOTALL | re.IGNORECASE)
                impacts = re.findall(r"<div class=['\"]risk-impact['\"]>(.*?)</div>", block_html, re.DOTALL | re.IGNORECASE)
                mitigations = re.findall(r"<strong>Mitigation:</strong> (.*?)</p>", block_html, re.DOTALL | re.IGNORECASE)

                for idx in range(min(len(titles), len(descs), len(impacts))):
                    mit_text = mitigations[idx] if idx < len(mitigations) else "Standard testing procedures."

                    risk_box_data = [
                        [Paragraph(f"<b>RISK: {clean_html_for_reportlab(titles[idx])}</b>", styles['CorpBody']),
                         Paragraph(f"<font color='#c8102e'><b>IMPACT: {clean_html_for_reportlab(impacts[idx]).upper()}</b></font>", styles['CorpTableText'])],
                        [Paragraph(f"<i>Description:</i> {clean_html_for_reportlab(descs[idx])}", styles['CorpTableText']), ""],
                        [Paragraph(f"<b>Mitigation:</b> {clean_html_for_reportlab(mit_text)}", styles['CorpTableText']), ""]
                    ]

                    t = Table(risk_box_data, colWidths=[5 * inch, 1.5 * inch])
                    t.setStyle(TableStyle([
                        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#fcf8f8')),
                        ('LINELEFT', (0, 0), (0, -1), 3, brand_red),
                        ('BOX', (0, 0), (-1, -1), 0.5, border_grey),
                        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                        ('SPAN', (0, 1), (1, 1)),
                        ('SPAN', (0, 2), (1, 2)),
                        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
                        ('TOPPADDING', (0, 0), (-1, -1), 4),
                    ]))
                    story.append(t)
                    story.append(Spacer(1, 8))
            except Exception as e:
                story.append(Paragraph(clean_html_for_reportlab(block_html), styles['CorpBody']))
                story.append(Spacer(1, 6))

        else:
            clean_text = clean_html_for_reportlab(block_html)
            if clean_text:
                story.append(Paragraph(clean_text, styles['CorpBody']))
                story.append(Spacer(1, 6))

    remaining_text = content[last_end:].strip()
    if remaining_text:
        clean_text = clean_html_for_reportlab(remaining_text)
        if clean_text:
            story.append(Paragraph(clean_text, styles['CorpBody']))
            story.append(Spacer(1, 6))


def build_enterprise_pdf(workspace_data, filepath):
    # Brand colors matching KaarTech (Red and Navy)
    brand_red = colors.HexColor('#c8102e')
    brand_navy = colors.HexColor('#0f1b29')
    text_dark = colors.HexColor('#1f2937')
    border_grey = colors.HexColor('#e5e7eb')

    doc = BaseDocTemplate(filepath, pagesize=A4, rightMargin=inch, leftMargin=inch, topMargin=inch, bottomMargin=inch)
    
    # Custom Page Template with clean red & navy headers and footers
    class CorporateTemplate(PageTemplate):
        def __init__(self, id, title):
            self.title = title
            frames = [Frame(inch, inch, A4[0] - 2*inch, A4[1] - 2*inch, id='normal')]
            super().__init__(id, frames=frames)

        def beforeDrawPage(self, canvas, doc):
            if doc.page == 1:
                return
            canvas.saveState()
            canvas.setFont('Helvetica-Bold', 8)
            canvas.setFillColor(brand_navy)
            canvas.drawString(inch, A4[1] - 0.55*inch, "KTERN ENTERPRISE AI WORKSPACE")
            canvas.setFont('Helvetica', 8)
            canvas.setFillColor(colors.HexColor('#6b7280'))
            canvas.drawRightString(A4[0] - inch, A4[1] - 0.55*inch, f"Document: {self.title[:40]}...")
            
            # Header Line (Red)
            canvas.setStrokeColor(brand_red)
            canvas.setLineWidth(1)
            canvas.line(inch, A4[1] - 0.6*inch, A4[0] - inch, A4[1] - 0.6*inch)
            
            # Footer Line (Grey)
            canvas.setStrokeColor(border_grey)
            canvas.line(inch, 0.7*inch, A4[0] - inch, 0.7*inch)
            
            # Footer text
            canvas.setFont('Helvetica', 8)
            canvas.drawString(inch, 0.5*inch, "Confidential - Powered by KaarTech Solutions")
            canvas.drawRightString(A4[0] - inch, 0.5*inch, f"Page {doc.page}")
            canvas.restoreState()

    doc.addPageTemplates([CorporateTemplate('Corporate', workspace_data.get('title', 'Enterprise Proposal'))])

    styles = getSampleStyleSheet()
    
    # Custom Corporate styles
    styles.add(ParagraphStyle(name='CorpCoverTitle', parent=styles['Heading1'], fontSize=24, leading=30, spaceAfter=15, alignment=1, textColor=brand_navy))
    styles.add(ParagraphStyle(name='CorpCoverClient', parent=styles['Normal'], fontSize=11, leading=15, spaceAfter=20, alignment=1, textColor=brand_red))
    styles.add(ParagraphStyle(name='CorpHeading', parent=styles['Heading1'], fontSize=14, leading=18, spaceBefore=15, spaceAfter=8, textColor=brand_navy, keepWithNext=True))
    styles.add(ParagraphStyle(name='CorpBody', parent=styles['BodyText'], fontSize=9.5, leading=13.5, spaceAfter=8, textColor=text_dark))
    styles.add(ParagraphStyle(name='CorpBullet', parent=styles['Normal'], fontSize=9.5, leading=13.5, leftIndent=15, spaceAfter=4, textColor=text_dark))
    styles.add(ParagraphStyle(name='CorpTableText', parent=styles['Normal'], fontSize=8.5, leading=11, textColor=text_dark))
    styles.add(ParagraphStyle(name='CorpTableHeader', parent=styles['Normal'], fontSize=8.5, leading=11, fontName='Helvetica-Bold', textColor=colors.whitesmoke))

    story = []

    # 1. Cover Page Flow
    story.append(Spacer(1, 1.5*inch))
    story.append(Paragraph("<font size='22' color='#c8102e'><i><b>K</b></i></font><font size='18' color='#0f1b29'><b>Tern</b></font>", styles['CorpCoverTitle']))
    story.append(Spacer(1, 0.3*inch))
    story.append(Paragraph(workspace_data.get('title', 'Enterprise Proposal').upper(), styles['CorpCoverTitle']))
    story.append(Spacer(1, 0.15*inch))
    story.append(Paragraph("Multi-Agent Enterprise Document Intelligence Platform", styles['CorpCoverClient']))
    story.append(Spacer(1, 1*inch))
    
    story.append(Paragraph(f"<b>PREPARED FOR:</b> {workspace_data.get('client', 'KaarTech Solutions')}", styles['CorpCoverClient']))
    story.append(Paragraph(f"<b>DATE:</b> {workspace_data.get('date', datetime.now().strftime('%B %d, %Y'))}", styles['CorpCoverClient']))
    story.append(Spacer(1, 0.5*inch))
    story.append(Paragraph("<font color='#6b7280'>CONFIDENTIAL &middot; FOR ENTERPRISE REVIEW ONLY</font>", styles['CorpCoverClient']))
    story.append(PageBreak())

    # 2. Add sections from workspace data
    for sec in workspace_data.get('sections', []):
        story.append(Paragraph(clean_html_for_reportlab(sec.get('title', 'Section')), styles['CorpHeading']))
        content = sec.get('content', '')
        
        parse_html_to_story(content, styles, story, brand_navy, brand_red, border_grey, A4)
        story.append(Spacer(1, 10))
        
    doc.build(story)





@app.route('/')
def index():
    return render_template('index.html')


@app.route('/generate', methods=['POST'])
def generate():
    if request.is_json:
        data = request.json
        topic = data.get('topic')
        requirements = data.get('requirements', '')
    else:
        topic = request.form.get('topic')
        requirements = request.form.get('requirements', '')
    
    if not topic:
        return jsonify({"error": "Topic is required"}), 400
        
    try:
        # Check if the requirements parameter contains a JSON payload from the workspace
        if requirements.strip().startswith('{') and requirements.strip().endswith('}'):
            try:
                workspace_data = json.loads(requirements)
                if isinstance(workspace_data, dict) and "sections" in workspace_data:
                    print(f"Generating Premium Enterprise PDF from workspace state: {topic}")
                    filename = f"KaarTech_Proposal_{uuid.uuid4().hex[:8]}.pdf"
                    filepath = os.path.join(PDF_DIR, filename)
                    
                    build_enterprise_pdf(workspace_data, filepath)
                    
                    pdf_url = url_for('download_pdf', filename=filename)
                    return jsonify({"success": True, "pdf_url": pdf_url})
            except Exception as json_err:
                print(f"Failed to compile workspace JSON directly: {json_err}. Falling back to AI call.")
                
        print(f"Generating Enterprise Document for: {topic}")
        # 1. Fetch JSON structured content from OpenRouter
        report_data = generate_report_data(topic, requirements)
        
        # 2. Build PDF Document
        filename = f"EnterpriseDoc_{uuid.uuid4().hex[:8]}.pdf"
        filepath = os.path.join(PDF_DIR, filename)
        
        build_enterprise_pdf(report_data, filepath)
        
        pdf_url = url_for('download_pdf', filename=filename)
        return jsonify({"success": True, "pdf_url": pdf_url})
        
    except Exception as e:
        print(f"Error during generation: {str(e)}")
        import traceback
        traceback.print_exc()
        return jsonify({"error": str(e)}), 500


@app.route('/upload', methods=['POST'])
def upload_file():
    if 'file' not in request.files:
        return jsonify({"error": "No file part"}), 400
    file = request.files['file']
    if file.filename == '':
        return jsonify({"error": "No selected file"}), 400
        
    filename = secure_filename(file.filename)
    ext = filename.rsplit('.', 1)[1].lower() if '.' in filename else ''
    
    extracted_text = ""
    try:
        if ext == 'pdf':
            with pdfplumber.open(file) as pdf:
                extracted_text = "\n".join([page.extract_text() for page in pdf.pages if page.extract_text()])
        elif ext in ['doc', 'docx']:
            doc = docx.Document(file)
            extracted_text = "\n".join([para.text for para in doc.paragraphs])
        elif ext == 'csv':
            df = pd.read_csv(file)
            extracted_text = df.to_string()
        elif ext == 'txt':
            extracted_text = file.read().decode('utf-8')
        else:
            return jsonify({"error": "Unsupported file format. Please upload PDF, DOCX, CSV, or TXT."}), 400
            
        return jsonify({
            "success": True, 
            "filename": filename,
            "extracted_text": extracted_text[:15000] # Limit size to prevent token blowup
        })
    except Exception as e:
        print(f"Error extracting file: {e}")
        return jsonify({"error": f"Failed to extract text: {str(e)}"}), 500

@app.route('/chat', methods=['POST'])
def chat():
    data = request.json
    message = data.get('message', '')
    history = data.get('history', [])
    doc_context = data.get('doc_context', {})
    
    if not message:
        return jsonify({"error": "Message is required"}), 400
        
    try:
        openrouter_key = os.getenv('OPENROUTER_API_KEY', '')
        if not openrouter_key:
            return jsonify({"reply": "KTern En offline assistant: Please set OPENROUTER_API_KEY."})
            
        prompt = f"""
        You are KTern En, the expert enterprise AI migration and systems consulting assistant by KaarTech.
        The user is discussing a proposal or document in the workspace.
        
        Current Document Title in Workspace: {doc_context.get('title', 'SAP ECC to S/4HANA Migration Proposal')}
        Current Document Sections: {json.dumps(doc_context.get('sections', []))}
        
        User's Message: {message}
        
        You must decide if the user's message is just a general question, OR if they want you to update/rewrite a specific section of the document, OR if they want you to insert/create a new section in the document (e.g. adding a diagram, timeline, list of benefits, or custom analysis).
        
        Output MUST be strict JSON:
        {{
          "reply": "Your conversational reply to the user (e.g. 'I will add the architecture diagram section now')",
          "action": "update_section" | "insert_section" | "none",
          "section_id": "the-id-of-the-section (e.g. 'architecture')",
          "section_title": "The Title of the Section (e.g. 'System Landscape Architecture')",
          "section_icon": "FontAwesome icon class (e.g. 'fa-diagram-project')",
          "content": "<p>The HTML content for the section. If the user asked for a diagram, you MUST include a valid Mermaid flowchart inside a <div class=\\\"mermaid\\\">...</div> tag.</p>"
        }}
        
        Do not use markdown blocks around the JSON. Provide clean, actionable advice.
        """
        
        headers = {
            "Authorization": f"Bearer {openrouter_key}",
            "Content-Type": "application/json",
            "HTTP-Referer": "http://localhost:5000",
            "X-Title": "KTern En Chatbot"
        }
        
        messages = [{"role": "system", "content": "You are KTern En, a professional SAP consulting AI. Output ONLY valid JSON."}]
        for h in history:
            role = "user" if h.get("sender") == "user" else "assistant"
            messages.append({"role": role, "content": h.get("text", "")})
        messages.append({"role": "user", "content": prompt})
        
        payload = {
            "model": AI_MODEL,
            "max_tokens": 2000,
            "messages": messages,
            "temperature": 0.5
        }
        
        r = requests.post("https://openrouter.ai/api/v1/chat/completions", headers=headers, json=payload, timeout=20)
        r.raise_for_status()
        reply_text = r.json()['choices'][0]['message']['content']
        
        reply_text = re.sub(r'^```json\s*', '', reply_text)
        reply_text = re.sub(r'^```\s*', '', reply_text)
        reply_text = re.sub(r'\s*```$', '', reply_text)
        
        reply_json = json.loads(reply_text)
        action = reply_json.get("action", "none")
        section_id = reply_json.get("section_id")
        section_title = reply_json.get("section_title")
        section_icon = reply_json.get("section_icon", "fa-file-lines")
        content = reply_json.get("content")
        
        update_section_data = None
        insert_section_data = None
        
        if action == "update_section" and section_id:
            update_section_data = {
                "id": section_id,
                "content": content
            }
            if section_title:
                update_section_data["title"] = section_title
            if section_icon:
                update_section_data["icon"] = section_icon
        elif action == "insert_section" and section_id:
            insert_section_data = {
                "id": section_id,
                "title": section_title or "New Section",
                "icon": section_icon or "fa-file-lines",
                "content": content
            }
            
        return jsonify({
            "success": True, 
            "reply": reply_json.get("reply", "Done."),
            "action": action,
            "update_section": update_section_data,
            "insert_section": insert_section_data
        })
        
    except Exception as e:
        print(f"Error in chat endpoint: {e}")
        import traceback
        traceback.print_exc()
        return jsonify({"reply": f"As your KTern consulting assistant, I've analyzed your query: '{message}'. Currently facing an error: {e}"})


@app.route('/assist', methods=['POST'])
def assist():
    data = request.json
    section_title = data.get('section_title', '')
    section_content = data.get('section_content', '')
    action = data.get('action', 'expand')
    doc_title = data.get('doc_title', 'SAP ECC to S/4HANA Migration Proposal')
    
    if not section_title:
        return jsonify({"error": "Section title is required"}), 400
        
    try:
        openrouter_key = os.getenv('OPENROUTER_API_KEY', '')
        if not openrouter_key:
            # Fallback offline templates
            if action == 'expand':
                expanded = section_content + f"<p>Additionally, our enterprise orchestration tools monitor the migration pathways in the target environments to automatically trigger safety fallbacks. We perform rigorous custom code checks, Z-table analysis, and data mapping validations. This strategy completely guarantees 100% data consistency, eliminates standard transaction locks, and facilitates absolute zero-friction cutover windows. Furthermore, business process alignment dashboards are deployed to give key stakeholders real-time visibility into active operational throughput during migration cycles.</p>"
                return jsonify({"success": True, "content": expanded})
            elif action == 'shorten':
                return jsonify({"success": True, "content": f"<p>Modernized enterprise platform with zero operational downtime. All processes are fully optimized and verified using real-time validation tools.</p>"})
            else:
                return jsonify({"success": True, "content": f"<p>Pursuant to organizational mandates, all system components will enforce strict standard configurations. Clean-core execution rules govern all standard custom extensions to secure structural integrity.</p>"})

        if action == 'expand':
            instruction = "Rewrite and greatly expand the following HTML content to be 3 to 4 times longer. Make it extremely detailed, comprehensive, high-density, and professional. Cover specific SAP processes, transaction codes, architectures, validation checklists, and expert recommendations where applicable. You MUST output well-styled HTML matching the input's format (e.g. keeping <ul>, <li>, <table>, or <p> tags but making the copy extremely rich and dense). Do not wrap inside a ```html block, return only the raw HTML code."
        elif action == 'shorten':
            instruction = "Condense and simplify the following HTML content to be highly concise, punchy, and clear while retaining the essential details. Return only raw HTML."
        else:
            instruction = "Rewrite the following HTML content using an extremely polished, high-end corporate executive vocabulary perfect for a C-suite presentation. Retain the same length and keep the HTML markup. Return only raw HTML."

        prompt = f"""
        You are an elite enterprise architect and lead systems consultant by KaarTech.
        
        Document Context: {doc_title}
        Section Title: {section_title}
        Current Section Content (HTML): {section_content}
        
        Action requested: {action.upper()}
        Instruction: {instruction}
        """

        headers = {
            "Authorization": f"Bearer {openrouter_key}",
            "Content-Type": "application/json",
            "HTTP-Referer": "http://localhost:5000",
            "X-Title": "KTern Section Assistant"
        }
        
        payload = {
            "model": AI_MODEL,
            "max_tokens": 2000,
            "messages": [
                {"role": "system", "content": "You are a professional HTML-generating consulting assistant. Output ONLY valid, raw HTML without any markdown code block wrap."},
                {"role": "user", "content": prompt}
            ],
            "temperature": 0.5
        }
        
        r = requests.post("https://openrouter.ai/api/v1/chat/completions", headers=headers, json=payload, timeout=30)
        r.raise_for_status()
        content = r.json()['choices'][0]['message']['content']
        
        # Clean up any potential markdown code blocks if returned
        content = re.sub(r'^```html\s*', '', content)
        content = re.sub(r'^```\s*', '', content)
        content = re.sub(r'\s*```$', '', content)
        
        return jsonify({"success": True, "content": content.strip()})
        
    except Exception as e:
        print(f"Error in assist endpoint: {e}")
        return jsonify({"error": str(e)}), 500



@app.route('/generate_doc', methods=['POST'])
def generate_doc():
    if request.is_json:
        data = request.json
        topic = data.get('topic')
        requirements = data.get('requirements', '')
    else:
        topic = request.form.get('topic')
        requirements = request.form.get('requirements', '')
        
    if not topic:
        return jsonify({"error": "Topic is required"}), 400
        
    try:
        print(f"Generating dynamic document content for topic: {topic}")
        report_data = generate_report_data(topic, requirements)
        return jsonify({"success": True, "data": report_data})
    except Exception as e:
        print(f"Error during document generation: {str(e)}")
        import traceback
        traceback.print_exc()
        return jsonify({"error": str(e)}), 500


@app.route('/download/<filename>')
def download_pdf(filename):
    return send_from_directory(PDF_DIR, filename)


if __name__ == '__main__':
    app.run(debug=True, port=int(os.getenv("PORT", 5000)))
